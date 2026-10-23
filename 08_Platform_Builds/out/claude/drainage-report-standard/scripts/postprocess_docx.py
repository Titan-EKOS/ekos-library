"""DOCX post-process and check for drafts written by assemble_report.py (Standard v0.3 chain).

Fixes (only with --fix, writes --out):
  - core properties: author / last_modified_by = firm preparer string; comments, keywords, subject cleared
  - DATE / TIME / PRINTDATE / SAVEDATE fields removed from headers and footers
  - template custom properties part (docProps/custom.xml) dropped
Checks (always; nothing here deletes report content):
  - unresolved {placeholders}, [FIGURE:..] / [TABLE:..] / [LIB:..] slots, rewrite tags   (kept and listed, D44)
  - interim wording (TBD, TBC, TO BE CONFIRMED, XX.X, "?" references)
  - tool / assistant names in text or properties
  - body paragraphs without a named style
  - tables: shading on body rows (D47 = no alternating shading), header row not set to repeat

Usage:
  "%DOCS_PY%" postprocess_docx.py --docx <draft.docx> [--fix --out <draft_pp.docx> --preparer "Psomas (JCO)"] [--json <r.json>]
Exit codes: 0 no findings, 1 findings listed, 2 usage error.
"""
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

PLACEHOLDER = re.compile(r"\{[A-Za-z_][A-Za-z0-9_ |.-]*\}")
SLOT = re.compile(r"\[(FIGURE|TABLE|LIB):[^\]]*\]")
REWRITE = re.compile("⟦REWRITE[^⟧]*⟧")
INTERIM = re.compile(r"\b(TBD|TBC|TO BE CONFIRMED|TO BE DETERMINED)\b|XX\.X|\b[A-Z]{1,3}\d*-\?|\(\?\)")
TOOLNAMES = re.compile(r"\b(Claude|ChatGPT|GPT-?\d|Gemini|Copilot|Codex|Anthropic|OpenAI|LLM|language model|AI[- ]generated)\b", re.I)
DATE_FIELD = re.compile(r"^\s*(DATE|TIME|PRINTDATE|SAVEDATE|CREATEDATE)\b", re.I)


def iter_paragraphs(doc):
    body = doc.element.body
    for p in body.iter(qn("w:p")):
        yield p


def p_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def remove_date_fields(part_element) -> int:
    removed = 0
    for fs in list(part_element.iter(qn("w:fldSimple"))):
        if DATE_FIELD.match(fs.get(qn("w:instr")) or ""):
            fs.getparent().remove(fs)
            removed += 1
    # complex fields: fldChar begin ... instrText ... separate ... result ... end, all inside runs of one paragraph
    for p in part_element.iter(qn("w:p")):
        runs = list(p.iter(qn("w:r")))
        i = 0
        while i < len(runs):
            fc = runs[i].find(qn("w:fldChar"))
            if fc is not None and fc.get(qn("w:fldCharType")) == "begin":
                depth, j, instr = 0, i, ""
                while j < len(runs):
                    f = runs[j].find(qn("w:fldChar"))
                    it = runs[j].find(qn("w:instrText"))
                    if it is not None:
                        instr += it.text or ""
                    if f is not None:
                        t = f.get(qn("w:fldCharType"))
                        depth += 1 if t == "begin" else (-1 if t == "end" else 0)
                        if depth == 0:
                            break
                    j += 1
                if DATE_FIELD.match(instr) and j < len(runs):
                    for r in runs[i:j + 1]:
                        r.getparent().remove(r)
                    removed += 1
                    runs = list(p.iter(qn("w:r")))
                    continue
                i = j + 1
            else:
                i += 1
    return removed


def table_findings(doc):
    out = []
    for n, tbl in enumerate(doc.element.body.iter(qn("w:tbl")), start=1):
        rows = tbl.findall(qn("w:tr"))
        if not rows:
            continue
        trpr = rows[0].find(qn("w:trPr"))
        if trpr is None or trpr.find(qn("w:tblHeader")) is None:
            out.append(f"table {n}: header row not set to repeat")
        shaded = 0
        for r in rows[1:]:
            for shd in r.iter(qn("w:shd")):
                fill = (shd.get(qn("w:fill")) or "auto").upper()
                if fill not in ("AUTO", "FFFFFF", "NONE"):
                    shaded += 1
                    break
        if shaded:
            out.append(f"table {n}: {shaded} body row(s) carry shading (D47: no alternating shading)")
        look = tbl.find(qn("w:tblPr"))
        if look is not None:
            tl = look.find(qn("w:tblLook"))
            if tl is not None and tl.get(qn("w:noHBand")) in ("0", "false"):
                out.append(f"table {n}: table style banding switched on (tblLook noHBand=0)")
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--docx", required=True, type=Path)
    ap.add_argument("--fix", action="store_true")
    ap.add_argument("--out", type=Path)
    ap.add_argument("--preparer", default="Psomas")
    ap.add_argument("--json", type=Path)
    a = ap.parse_args(argv)
    if a.fix and not a.out:
        print("--fix needs --out (issued drafts are never overwritten)", file=sys.stderr)
        return 2
    doc = Document(str(a.docx))
    res = {"docx": str(a.docx), "placeholders": [], "slots": [], "rewrite_tags": 0, "interim": [], "tool_names": [],
           "unstyled_paragraphs": 0, "tables": [], "fixed": {}}

    for p in iter_paragraphs(doc):
        t = p_text(p)
        if not t.strip():
            continue
        res["placeholders"] += PLACEHOLDER.findall(t)
        res["slots"] += [m.group(0) for m in SLOT.finditer(t)]
        res["rewrite_tags"] += len(REWRITE.findall(t))
        res["interim"] += [m.group(0) for m in INTERIM.finditer(t)]
        res["tool_names"] += [m.group(0) for m in TOOLNAMES.finditer(t)]
        ppr = p.find(qn("w:pPr"))
        if ppr is None or ppr.find(qn("w:pStyle")) is None:
            res["unstyled_paragraphs"] += 1
    cp = doc.core_properties
    for fld in ("author", "last_modified_by", "comments", "keywords", "subject", "title", "category"):
        v = getattr(cp, fld) or ""
        if TOOLNAMES.search(v):
            res["tool_names"].append(f"core property {fld}: {v}")
    res["tables"] = table_findings(doc)

    if a.fix:
        cp.author = a.preparer
        cp.last_modified_by = a.preparer
        cp.comments = ""
        cp.keywords = ""
        cp.subject = ""
        n = 0
        for sec in doc.sections:
            for hf in (sec.header, sec.footer, sec.first_page_header, sec.first_page_footer,
                       sec.even_page_header, sec.even_page_footer):
                try:
                    n += remove_date_fields(hf._element)
                except Exception:
                    pass
        res["fixed"]["date_fields_removed"] = n
        pkg = doc.part.package
        for rel_id, rel in list(pkg.rels.items()):
            if rel.reltype.endswith("/custom-properties"):
                del pkg.rels[rel_id]
                res["fixed"]["custom_properties_dropped"] = True
        res["fixed"]["author"] = a.preparer
        doc.save(str(a.out))
        res["out"] = str(a.out)

    findings = (len(res["placeholders"]) + len(res["slots"]) + res["rewrite_tags"] + len(res["interim"])
                + len(res["tool_names"]) + res["unstyled_paragraphs"] + len(res["tables"]))
    if a.json:
        a.json.write_text(json.dumps(res, indent=1), encoding="utf-8")
    print(f"placeholders {len(res['placeholders'])} | slots {len(res['slots'])} | rewrite tags {res['rewrite_tags']} | "
          f"interim {len(res['interim'])} | tool names {len(res['tool_names'])} | unstyled {res['unstyled_paragraphs']} | "
          f"table findings {len(res['tables'])}" + (f" | fixed {res['fixed']}" if a.fix else ""))
    return 1 if findings else 0


if __name__ == "__main__":
    sys.exit(main())
