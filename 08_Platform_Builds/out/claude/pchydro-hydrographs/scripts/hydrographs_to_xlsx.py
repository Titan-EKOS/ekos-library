#!/usr/bin/env python3
"""
hydrographs_to_xlsx.py -- Phase 4 (export QC) of pchydro-hydrographs v2.

Combines PC-Hydro hydrograph exports into one QC workbook and, per decision D19,
keeps two series families apart:

  Native         the export exactly as PC-Hydro wrote it (record copy, never edited)
  Peak_Adjusted  a separate copy for routing: the maximum ordinate is replaced by the
                 program-reported peak (when supplied and higher than the sampled max)
                 and a nonzero tail is closed with one extra step at 0 cfs. Every added
                 volume is logged.

The Peaks sheet reports BOTH the sampled maximum and the program-reported peak.
Report tables use the program-reported peak (native output is the reference for new
reports, 0.0 cfs tolerance). For rebuilds of an old report, pass --targets to compare
against the old report's peaks.

Accepted inputs (any mix):
  a) Native PC-Hydro "Generate Hydrograph" CSV: header lines containing
     "Concentration Point", then per event a "Return Period:" line, a
     "Peak Discharge:" line and time,flow rows beneath. The reported peak is read
     from the "Peak Discharge:" line.
  b) Plain table CSV: first column time, other columns named cfs series. Lines
     starting with '#' are ignored. Reported peaks come from --peaks.

Optional side files:
  --peaks   CSV with columns: series, reported_peak_cfs [, equation_volume_acft]
            (overrides or fills the reported peak; equation volume = the program's
            runoff-volume equation result, reported separately, never used as a gate)
  --targets CSV with columns: series, target_peak_cfs (old-report peaks for rebuilds)

No volume-equality gate is applied: the county manual treats the integrated
hydrograph as an approximation of the equation volume (EKOS-LOCAL-007).

Usage (EKOS, GIS venv python by full path):
  D:\\EKOS\\venvs\\gis\\Scripts\\python.exe hydrographs_to_xlsx.py --out CP_hydrographs.xlsx ^
      --adjusted-csv CP_routing_series.csv --native-csv CP_native_series.csv CP-101.csv CP-102.csv
Container:
  python3 hydrographs_to_xlsx.py --out combined.xlsx --peaks peaks.csv series1.csv
"""
import argparse
import csv
import os
import re
import sys

CF_PER_ACFT = 43560.0
STEP_TOL_S = 0.5          # tolerance on the 60 s spacing check
PEAK_TOL_CFS = 0.005      # below this the sampled max is treated as equal to the reported peak

RE_CP = re.compile(r"concentration\s*point\s*[:,;\t ]*(.*)", re.I)
RE_RP = re.compile(r"return\s*period\s*[:,;\t ]*([\d.]+)", re.I)
RE_PK = re.compile(r"peak\s*discharge\s*[:,;\t ]*(\d[\d,]*\.?\d*)", re.I)


def num(x):
    try:
        return float(str(x).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


class Series:
    def __init__(self, name, source, cp="", rp=""):
        self.name = name
        self.source = source
        self.cp = cp
        self.rp = rp
        self.t = []          # minutes
        self.q = []          # cfs
        self.reported_peak = None
        self.equation_volume = None
        self.target_peak = None


def clean_name(s):
    s = re.sub(r"[^A-Za-z0-9_.-]+", "_", s.strip()).strip("_")
    return s or "SERIES"


def parse_native(path, lines):
    """Parse a native PC-Hydro export into one Series per return period."""
    out, cp, cur = [], "", None
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        m = RE_CP.search(line)
        if m and not RE_RP.search(line):
            cp = m.group(1).strip().strip(",").strip('"')
            continue
        m = RE_RP.search(line)
        if m:
            rp = m.group(1).rstrip(".")
            rp_txt = rp[:-2] if rp.endswith(".0") else rp
            cur = Series(f"{clean_name(cp or os.path.splitext(os.path.basename(path))[0])}_{rp_txt}yr",
                         os.path.basename(path), cp, rp_txt)
            out.append(cur)
            continue
        m = RE_PK.search(line)
        if m and cur is not None:
            cur.reported_peak = num(m.group(1))
            continue
        if cur is None:
            continue
        cells = [c for c in re.split(r"[,\t;]", line) if c.strip() != ""]
        if len(cells) >= 2:
            tv, qv = num(cells[0]), num(cells[1])
            if tv is not None and qv is not None:
                cur.t.append(tv)
                cur.q.append(qv)
    return [s for s in out if s.q]


def parse_table(path, lines):
    rows = list(csv.reader([l for l in lines if l.strip() and not l.lstrip().startswith("#")]))
    if not rows:
        return []
    hdr, data = rows[0], rows[1:]
    out = []
    for j in range(1, len(hdr)):
        s = Series(clean_name(hdr[j]), os.path.basename(path))
        for r in data:
            if len(r) <= j:
                continue
            tv, qv = num(r[0]), num(r[j])
            if tv is not None and qv is not None:
                s.t.append(tv)
                s.q.append(qv)
        if s.q:
            out.append(s)
    return out


def load(path, time_units):
    with open(path, errors="ignore") as fh:
        lines = fh.read().splitlines()
    native = any(RE_RP.search(l) for l in lines)
    series = parse_native(path, lines) if native else parse_table(path, lines)
    for s in series:
        to_minutes(s, time_units)
    return series, native


def median_step(t):
    d = sorted(b - a for a, b in zip(t, t[1:]))
    return d[len(d) // 2] if d else 0.0


def to_minutes(s, units):
    if units == "hr" or (units == "auto" and 0 < median_step(s.t) < 0.1):
        s.t = [round(v * 60.0, 4) for v in s.t]


def step_check(t):
    """Return (median step in s, True if every step is 60 s within tolerance)."""
    if len(t) < 2:
        return 0.0, False
    steps_s = [(b - a) * 60.0 for a, b in zip(t, t[1:])]
    med = sorted(steps_s)[len(steps_s) // 2]
    ok = all(abs(x - 60.0) <= STEP_TOL_S for x in steps_s)
    return med, ok


def volume_acft(t, q):
    """Trapezoidal integral of q (cfs) over t (min), in ac-ft."""
    v = 0.0
    for i in range(1, len(t)):
        v += 0.5 * (q[i] + q[i - 1]) * (t[i] - t[i - 1]) * 60.0
    return v / CF_PER_ACFT


def adjust(s):
    """Build the peak-adjusted copy. Returns (t, q, info dict)."""
    t, q = list(s.t), list(s.q)
    info = {"sampled_max": max(q), "sampled_time": t[q.index(max(q))],
            "peak_added_acft": 0.0, "tail_added_acft": 0.0, "notes": []}
    v0 = volume_acft(t, q)
    rp = s.reported_peak
    if rp is None:
        info["notes"].append("no reported peak supplied; maximum ordinate not adjusted")
    elif rp - info["sampled_max"] > PEAK_TOL_CFS:
        i = q.index(info["sampled_max"])
        q[i] = rp
        info["peak_added_acft"] = volume_acft(t, q) - v0
        info["notes"].append(f"max ordinate at t={t[i]:g} min raised {info['sampled_max']:.3f} -> {rp:.3f} cfs")
    elif info["sampled_max"] - rp > PEAK_TOL_CFS:
        info["notes"].append(f"WARN reported peak {rp:.3f} below sampled max {info['sampled_max']:.3f} cfs; "
                             "not adjusted, check the export and the printout")
    else:
        info["notes"].append("sampled max equals reported peak; no peak adjustment")
    if q[-1] > 0:
        dt = (t[-1] - t[-2]) if len(t) > 1 else 1.0
        v1 = volume_acft(t, q)
        t.append(t[-1] + dt)
        q.append(0.0)
        info["tail_added_acft"] = volume_acft(t, q) - v1
        info["notes"].append(f"nonzero tail {q[-2]:.3f} cfs closed to 0 at t={t[-1]:g} min")
    return t, q, info


def read_side(path, key):
    out = {}
    if not path:
        return out
    with open(path, newline="", errors="ignore") as fh:
        for r in csv.DictReader(fh):
            r = {k.strip().lower(): v for k, v in r.items() if k}
            name = clean_name(r.get("series", ""))
            if name:
                out[name] = {k: num(v) for k, v in r.items() if k != "series"}
    return out


def write_series_csv(path, series_list, use_adjusted):
    """Write wide CSV(s) (Time_min + one cfs column per series) for build_dss.py.

    Series are grouped by time grid (start, step). Within a group, a series shorter
    than the longest is padded with trailing 0 cfs (valid because its tail ends at 0
    in the adjusted copy; in the native copy the padding is flagged in the header).
    Series on a different grid go to a separate file (<stem>_gridN.csv); ordinates are
    never interpolated or zero-filled inside a hydrograph.
    """
    groups = {}
    for s in series_list:
        t, q = (s.adj_t, s.adj_q) if use_adjusted else (s.t, s.q)
        step = round(median_step(t), 4)
        groups.setdefault((round(t[0], 4), step), []).append(
            (s.name + ("_ADJ" if use_adjusted else ""), t, q))
    stem, ext = os.path.splitext(path)
    written = []
    for gi, ((t0, step), cols) in enumerate(sorted(groups.items())):
        out = path if len(groups) == 1 else f"{stem}_grid{gi + 1}{ext or '.csv'}"
        n = max(len(t) for _, t, _ in cols)
        times = [round(t0 + i * step, 4) for i in range(n)]
        padded = [c for c, t, _ in cols if len(t) < n]
        with open(out, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["# " + ("peak-adjusted routing copy (D19)" if use_adjusted
                                else "native PC-Hydro export (record)")
                        + f"; step {step:g} min"])
            if padded:
                w.writerow(["# trailing 0 cfs padding added to: " + " ".join(padded)])
            w.writerow(["Time_min"] + [c for c, _, _ in cols])
            for i, tm in enumerate(times):
                w.writerow([f"{tm:g}"] + [f"{q[i]:.4f}" if i < len(q) else "0.0000" for _, _, q in cols])
        written.append(out)
        print("WROTE", out)
    if len(groups) > 1:
        print(f"WARN series on {len(groups)} different time grids; one CSV per grid. "
              "Re-export at 1 min so every series shares one grid.", file=sys.stderr)
    return written


def style_header(ws, font, fill):
    for c in ws[1]:
        if c.value is not None:
            c.font = font
            c.fill = fill


def put_block(ws, col, title, t, q):
    ws.cell(row=1, column=col, value="Time_min")
    ws.cell(row=1, column=col + 1, value=title)
    for i, (a, b) in enumerate(zip(t, q)):
        ws.cell(row=2 + i, column=col, value=a)
        ws.cell(row=2 + i, column=col + 1, value=b)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", required=True, help="output .xlsx")
    ap.add_argument("--peaks", help="CSV: series, reported_peak_cfs [, equation_volume_acft]")
    ap.add_argument("--targets", help="CSV: series, target_peak_cfs (old-report peaks, rebuilds only)")
    ap.add_argument("--target-tol", type=float, default=0.0, help="cfs tolerance vs --targets (default 0.0)")
    ap.add_argument("--time-units", choices=["auto", "min", "hr"], default="auto")
    ap.add_argument("--adjusted-csv", help="write the peak-adjusted routing series CSV (input to build_dss.py)")
    ap.add_argument("--native-csv", help="write the native series CSV (record)")
    ap.add_argument("csvs", nargs="+")
    a = ap.parse_args(argv)

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    peaks_side = read_side(a.peaks, "reported_peak_cfs")
    target_side = read_side(a.targets, "target_peak_cfs")

    allser = []
    for p in a.csvs:
        ser, native = load(p, a.time_units)
        if not ser:
            print(f"WARN no series parsed from {p}", file=sys.stderr)
        allser.extend(ser)
    if not allser:
        sys.exit("No hydrograph series found.")

    seen = {}
    for s in allser:                      # keep names unique
        if s.name in seen:
            seen[s.name] += 1
            s.name = f"{s.name}_{seen[s.name]}"
        else:
            seen[s.name] = 1
        side = peaks_side.get(s.name, {})
        if side.get("reported_peak_cfs") is not None:
            s.reported_peak = side["reported_peak_cfs"]
        s.equation_volume = side.get("equation_volume_acft")
        s.target_peak = target_side.get(s.name, {}).get("target_peak_cfs")
        s.adj_t, s.adj_q, s.info = adjust(s)

    wb = openpyxl.Workbook()
    hfont, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")

    pk = wb.active
    pk.title = "Peaks"
    pk_cols = ["series", "source_file", "concentration_point", "return_period_yr",
               "step_s", "step_60s_ok", "sampled_max_cfs", "sampled_peak_time_min",
               "reported_peak_cfs", "reported_minus_sampled_cfs", "report_table_value_cfs",
               "integrated_volume_native_acft", "peak_adjust_added_acft", "tail_close_added_acft",
               "integrated_volume_adjusted_acft", "equation_volume_acft",
               "target_peak_cfs", "target_diff_cfs", "target_status"]
    pk.append(pk_cols)
    style_header(pk, hfont, hfill)

    log = wb.create_sheet("Adjustment_Log")
    log.append(["series", "item", "detail"])
    style_header(log, hfont, hfill)

    nat = wb.create_sheet("Native")
    adj = wb.create_sheet("Peak_Adjusted")

    flags = 0
    col = 1
    for s in allser:
        step_s, step_ok = step_check(s.t)
        rp = s.reported_peak
        smax = s.info["sampled_max"]
        report_val = round(rp, 2) if rp is not None else None
        tdiff = tstat = None
        if s.target_peak is not None:
            basis = rp if rp is not None else smax
            tdiff = round(basis - s.target_peak, 3)
            tstat = "MATCH" if abs(tdiff) <= a.target_tol + 1e-9 else "MISMATCH"
        pk.append([s.name, s.source, s.cp, s.rp,
                   round(step_s, 2), "YES" if step_ok else "NO",
                   round(smax, 3), s.info["sampled_time"],
                   rp, round(rp - smax, 3) if rp is not None else None,
                   report_val if report_val is not None else "NO REPORTED PEAK",
                   round(volume_acft(s.t, s.q), 4),
                   round(s.info["peak_added_acft"], 5), round(s.info["tail_added_acft"], 5),
                   round(volume_acft(s.adj_t, s.adj_q), 4), s.equation_volume,
                   s.target_peak, tdiff, tstat])
        if not step_ok:
            log.append([s.name, "STEP", f"sample spacing not 60 s (median {step_s:.2f} s); re-export at 1 min"])
            flags += 1
        if rp is None:
            log.append([s.name, "PEAK", "no program-reported peak; report table value missing"])
            flags += 1
        for n in s.info["notes"]:
            log.append([s.name, "ADJUST", n])
            if n.startswith("WARN"):
                flags += 1
        if tstat == "MISMATCH":
            log.append([s.name, "TARGET", f"differs from old-report peak by {tdiff} cfs"])
            flags += 1
        put_block(nat, col, s.name, s.t, s.q)
        put_block(adj, col, s.name + "_ADJ", s.adj_t, s.adj_q)
        col += 3
    style_header(nat, hfont, hfill)
    style_header(adj, hfont, hfill)

    wb.save(a.out)
    print("WROTE", a.out, f"({len(allser)} series, {flags} flag(s); see Adjustment_Log)")
    if a.adjusted_csv:
        write_series_csv(a.adjusted_csv, allser, True)
    if a.native_csv:
        write_series_csv(a.native_csv, allser, False)
    return 0 if flags == 0 else 2


if __name__ == "__main__":
    sys.exit(main())
